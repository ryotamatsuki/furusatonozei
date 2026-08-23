#!/usr/bin/env python3
"""Independently reconcile every ordinary-grant field against MIC workbooks.

This validator intentionally does not use ``ordinary_grant.parse_source`` so
that the build parser and the audit parser do not share row-extraction logic.
"""
from __future__ import annotations

import hashlib
import io
import json
import math
import re
import sys
import unicodedata
from decimal import Decimal
from pathlib import Path
from urllib.request import Request, urlopen

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
PROCESSED = ROOT / "data" / "processed" / "furusato_data.json"
MANIFEST = ROOT / "data" / "ordinary_grant_manifest.json"
RAW_DIR = ROOT / "data" / "raw"
REPORT = ROOT / "data" / "processed" / "ordinary_grant_reconciliation_report.json"
SPECIAL_WARD_CODES = {f"131{i:02d}" for i in range(1, 24)}
FIELDS_PER_RECORD = 10


def normalize(value) -> str | None:
    if value is None:
        return None
    text = unicodedata.normalize("NFKC", str(value)).replace("\u3000", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text or None


def col_number(label: str) -> int:
    value = 0
    for char in label.upper():
        value = value * 26 + ord(char) - 64
    return value


def source_bytes(source: dict) -> bytes:
    path = RAW_DIR / source["file"]
    if path.exists():
        data = path.read_bytes()
    else:
        path.parent.mkdir(parents=True, exist_ok=True)
        request = Request(source["url"], headers={"User-Agent": "furusato-independent-audit/1.0"})
        with urlopen(request, timeout=120) as response:
            data = response.read()
        path.write_bytes(data)
    actual = hashlib.sha256(data).hexdigest()
    if actual != source["sha256"].lower():
        raise RuntimeError(f"SHA mismatch {source['file']}: expected {source['sha256']}, got {actual}")
    return data


def independent_source_rows(source: dict, records: list[dict]) -> dict[str, dict]:
    canonical = {(r["prefecture"], r["municipality"]): r["municipality_code"] for r in records}
    if len(canonical) != len(records):
        raise RuntimeError("canonical municipality names are not unique")
    workbook = load_workbook(io.BytesIO(source_bytes(source)), read_only=True, data_only=True)
    if source["sheet"] not in workbook.sheetnames:
        raise RuntimeError(f"sheet missing: {source['sheet']} / {workbook.sheetnames}")
    ws = workbook[source["sheet"]]
    p = col_number(source["columns"]["prefecture"]) - 1
    n = col_number(source["columns"]["municipality"]) - 1
    a = col_number(source["columns"]["amount_thousand_yen"]) - 1
    multiplier = Decimal(str(source.get("unit_multiplier", 1000)))
    current_pref = None
    rows: dict[str, dict] = {}
    for row_no, values in enumerate(ws.iter_rows(values_only=True), start=1):
        if row_no < int(source["row_start"]) or len(values) <= max(p, n, a):
            continue
        maybe_pref = normalize(values[p])
        if maybe_pref:
            current_pref = maybe_pref
        name = normalize(values[n])
        if not current_pref or not name:
            continue
        code = canonical.get((current_pref, name))
        if not code:
            continue
        if code in SPECIAL_WARD_CODES:
            raise RuntimeError(f"official ordinary-grant workbook unexpectedly includes special ward {code}")
        if code in rows:
            raise RuntimeError(f"duplicate official ordinary-grant row for {code}")
        raw_amount = values[a]
        if raw_amount is None or (isinstance(raw_amount, str) and not raw_amount.strip()):
            raise RuntimeError(f"blank ordinary-grant amount at row {row_no}")
        amount = Decimal(str(raw_amount).replace(",", "").strip()) * multiplier
        if amount != amount.to_integral_value() or amount < 0:
            raise RuntimeError(f"invalid ordinary-grant amount at row {row_no}: {raw_amount!r}")
        rows[code] = {"source_row": row_no, "amount": int(amount)}
    return rows


def same_number(actual, expected, tolerance=1e-5) -> bool:
    if actual is None or expected is None:
        return actual is expected
    return math.isclose(float(actual), float(expected), rel_tol=0, abs_tol=tolerance)


def audit() -> tuple[list[dict], dict]:
    normalized = json.loads(PROCESSED.read_text(encoding="utf-8"))
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    mismatches: list[dict] = []
    yearly = {}
    full_records = 0
    full_official_rows = 0
    full_wards = 0
    full_comparisons = 0

    for receipt_year, bucket in normalized["years"].items():
        records = bucket["records"]
        assessment_year = int(bucket["source"]["tax_assessment_fiscal_year"])
        source = manifest["sources"][str(assessment_year)]
        official = independent_source_rows(source, records)
        nonward = {r["municipality_code"] for r in records if r["municipality_code"] not in SPECIAL_WARD_CODES}
        wards = {r["municipality_code"] for r in records if r["municipality_code"] in SPECIAL_WARD_CODES}
        if set(official) != nonward:
            mismatches.append({
                "receipt_year": int(receipt_year),
                "field": "official_key_set",
                "missing": sorted(nonward - set(official)),
                "extra": sorted(set(official) - nonward),
            })
        if len(official) != int(manifest["expected_ordinary_municipality_count"]):
            mismatches.append({"receipt_year": int(receipt_year), "field": "official_row_count", "actual": len(official)})
        if len(wards) != int(manifest["expected_special_ward_count"]):
            mismatches.append({"receipt_year": int(receipt_year), "field": "special_ward_count", "actual": len(wards)})

        for record in records:
            code = record["municipality_code"]
            if code in SPECIAL_WARD_CODES:
                expected = {
                    "ordinary_grant_fiscal_year": assessment_year,
                    "ordinary_grant_amount": None,
                    "ordinary_grant_status": "special_ward_na",
                    "ordinary_grant_source_stage": source["stage"],
                    "ordinary_grant_source_stage_label": source["stage_label"],
                    "ordinary_grant_source_row": None,
                    "ordinary_grant_source_url": source["url"],
                    "ordinary_grant_source_sha256": source["sha256"],
                    "ordinary_grant_estimate": None,
                    "balance_with_ordinary_grant_estimate": None,
                }
            else:
                row = official.get(code)
                if row is None:
                    continue
                reference = float(record["municipal_tax_deduction"]) * 0.75
                estimate = min(reference, row["amount"])
                expected = {
                    "ordinary_grant_fiscal_year": assessment_year,
                    "ordinary_grant_amount": row["amount"],
                    "ordinary_grant_status": "recipient" if row["amount"] > 0 else "non_recipient",
                    "ordinary_grant_source_stage": source["stage"],
                    "ordinary_grant_source_stage_label": source["stage_label"],
                    "ordinary_grant_source_row": row["source_row"],
                    "ordinary_grant_source_url": source["url"],
                    "ordinary_grant_source_sha256": source["sha256"],
                    "ordinary_grant_estimate": estimate,
                    "balance_with_ordinary_grant_estimate": float(record["balance_before_tax_adjustment"]) + estimate,
                }
            for field, expected_value in expected.items():
                actual = record.get(field)
                full_comparisons += 1
                if isinstance(expected_value, (int, float)) and not isinstance(expected_value, bool):
                    ok = same_number(actual, expected_value)
                else:
                    ok = actual == expected_value
                if not ok:
                    mismatches.append({
                        "receipt_year": int(receipt_year),
                        "municipality_code": code,
                        "municipality": record["prefecture"] + record["municipality"],
                        "field": field,
                        "expected": expected_value,
                        "actual": actual,
                    })

        yearly[receipt_year] = {
            "processed_records": len(records),
            "official_ordinary_grant_rows": len(official),
            "special_ward_out_of_scope_records": len(wards),
            "field_comparisons": len(records) * FIELDS_PER_RECORD,
            "mismatches": sum(1 for m in mismatches if m.get("receipt_year") == int(receipt_year)),
            "ordinary_grant_fiscal_year": assessment_year,
            "ordinary_grant_file": source["file"],
            "ordinary_grant_sha256": source["sha256"],
        }
        full_records += len(records)
        full_official_rows += len(official)
        full_wards += len(wards)

    requested_years = [str(year) for year in range(2020, 2025)]
    requested = {
        "receipt_years": requested_years,
        "processed_records": sum(yearly[y]["processed_records"] for y in requested_years),
        "official_ordinary_grant_rows": sum(yearly[y]["official_ordinary_grant_rows"] for y in requested_years),
        "special_ward_out_of_scope_records": sum(yearly[y]["special_ward_out_of_scope_records"] for y in requested_years),
        "field_comparisons": sum(yearly[y]["field_comparisons"] for y in requested_years),
        "mismatches": sum(yearly[y]["mismatches"] for y in requested_years),
    }
    summary = {
        "yearly": yearly,
        "requested_2020_2024": requested,
        "full_period": {
            "processed_records": full_records,
            "official_ordinary_grant_rows": full_official_rows,
            "special_ward_out_of_scope_records": full_wards,
            "field_comparisons": full_comparisons,
            "mismatches": len(mismatches),
        },
    }
    return mismatches, summary


def main() -> None:
    mismatches, summary = audit()
    REPORT.parent.mkdir(parents=True, exist_ok=True)
    REPORT.write_text(json.dumps(mismatches, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    requested = summary["requested_2020_2024"]
    full = summary["full_period"]
    print("ORDINARY GRANT RECONCILIATION " + ("OK" if not mismatches else "FAILED"))
    print(f"requested 2020-2024 processed records: {requested['processed_records']}")
    print(f"requested official ordinary-grant rows: {requested['official_ordinary_grant_rows']}")
    print(f"requested special-ward out-of-scope records: {requested['special_ward_out_of_scope_records']}")
    print(f"requested record-field comparisons: {requested['field_comparisons']}")
    print(f"requested mismatches: {requested['mismatches']}")
    print(f"full processed records: {full['processed_records']}")
    print(f"full official ordinary-grant rows: {full['official_ordinary_grant_rows']}")
    print(f"full special-ward out-of-scope records: {full['special_ward_out_of_scope_records']}")
    print(f"full record-field comparisons: {full['field_comparisons']}")
    print(f"full mismatches: {full['mismatches']}")
    if mismatches:
        print(json.dumps(mismatches[:20], ensure_ascii=False, indent=2), file=sys.stderr)
        raise SystemExit(1)


if __name__ == "__main__":
    main()
