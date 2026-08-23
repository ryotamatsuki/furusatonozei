#!/usr/bin/env python3
"""Parse and apply official municipality ordinary local allocation tax data.

The official ordinary-grant amount is NOT treated as a causal reimbursement
from furusato nozei.  It is used only as a hard upper bound for a conservative
simple estimate:

    min(municipal tax deduction * 75%, total ordinary grant decision amount)

Tokyo's 23 special wards are outside municipality-unit ordinary-grant
estimation because Tokyo and the wards are treated together for local
allocation tax purposes.
"""
from __future__ import annotations

import hashlib
import io
import math
import re
import unicodedata
from decimal import Decimal, InvalidOperation
from pathlib import Path
from urllib.request import Request, urlopen

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
RAW_DIR = ROOT / "data" / "raw"
MAX_AMOUNT = 10**15
SPECIAL_WARD_CODES = {f"131{i:02d}" for i in range(1, 24)}


def fail(message: str) -> "NoReturn":
    raise RuntimeError(message)


def normalize_text(value) -> str | None:
    if value is None:
        return None
    text = unicodedata.normalize("NFKC", str(value)).replace("\u3000", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text or None


def excel_col_number(value: str) -> int:
    result = 0
    for char in value.upper():
        if not "A" <= char <= "Z":
            fail(f"invalid Excel column: {value}")
        result = result * 26 + ord(char) - ord("A") + 1
    return result


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def download_or_read(source: dict, *, no_download: bool) -> bytes:
    path = RAW_DIR / source["file"]
    data = path.read_bytes() if path.exists() else None
    if data is None and no_download:
        fail(f"missing cached ordinary-grant source: {path}")
    if data is None:
        path.parent.mkdir(parents=True, exist_ok=True)
        request = Request(
            source["url"],
            headers={"User-Agent": "furusato-data-integrity/3.0 (+https://github.com/ryotamatsuki/furusatonozei)"},
        )
        with urlopen(request, timeout=120) as response:
            data = response.read()
        path.write_bytes(data)
    actual = sha256_bytes(data)
    expected = source["sha256"].lower()
    if actual != expected:
        fail(
            f"ordinary-grant SHA-256 mismatch for {source['url']}: "
            f"expected {expected}, got {actual}"
        )
    return data


def parse_thousand_yen(value, *, row: int, multiplier: int) -> int:
    if value is None or (isinstance(value, str) and not value.strip()):
        fail(f"ordinary-grant blank amount at source row {row}")
    if isinstance(value, bool):
        fail(f"ordinary-grant boolean amount at source row {row}")
    try:
        dec = Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError) as exc:
        fail(f"ordinary-grant non-numeric amount {value!r} at source row {row}: {exc}")
    if not dec.is_finite() or dec < 0:
        fail(f"ordinary-grant invalid amount {value!r} at source row {row}")
    yen = dec * multiplier
    if yen != yen.to_integral_value():
        fail(f"ordinary-grant amount does not convert to integer yen at source row {row}: {value!r}")
    amount = int(yen)
    if amount > MAX_AMOUNT:
        fail(f"ordinary-grant implausibly large amount at source row {row}: {amount}")
    return amount


def is_special_ward(code5: str) -> bool:
    return code5 in SPECIAL_WARD_CODES


def parse_source(source: dict, canonical_by_name: dict[tuple[str, str], str], *, no_download: bool) -> dict[str, dict]:
    data = download_or_read(source, no_download=no_download)
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    if source["sheet"] not in workbook.sheetnames:
        fail(f"ordinary-grant sheet {source['sheet']!r} not found; available={workbook.sheetnames}")
    sheet = workbook[source["sheet"]]
    columns = source["columns"]
    pref_index = excel_col_number(columns["prefecture"]) - 1
    name_index = excel_col_number(columns["municipality"]) - 1
    amount_index = excel_col_number(columns["amount_thousand_yen"]) - 1
    start = int(source["row_start"])
    multiplier = int(source.get("unit_multiplier", 1000))

    rows: dict[str, dict] = {}
    current_pref: str | None = None
    duplicate_codes: list[str] = []
    for source_row, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if source_row < start:
            continue
        if len(row) <= max(pref_index, name_index, amount_index):
            continue
        raw_pref = normalize_text(row[pref_index])
        if raw_pref:
            current_pref = raw_pref
        name = normalize_text(row[name_index])
        if not current_pref or not name:
            continue
        code5 = canonical_by_name.get((current_pref, name))
        if not code5:
            # Aggregate/footer rows are deliberately ignored. Any real
            # municipality omission is detected by the exact count/missing
            # assertions below.
            continue
        if is_special_ward(code5):
            fail(f"ordinary-grant source unexpectedly contains Tokyo special ward {code5} {name}")
        if code5 in rows:
            duplicate_codes.append(code5)
            continue
        rows[code5] = {
            "source_row": source_row,
            "prefecture": current_pref,
            "municipality": name,
            "amount": parse_thousand_yen(row[amount_index], row=source_row, multiplier=multiplier),
        }
    if duplicate_codes:
        fail(f"ordinary-grant duplicate municipality rows: {duplicate_codes[:5]}")
    return rows


def enrich_records(records: list[dict], source: dict, *, no_download: bool, expected_count: int = 1718, expected_wards: int = 23) -> dict:
    canonical_by_name: dict[tuple[str, str], str] = {}
    nonward_codes: set[str] = set()
    ward_codes: set[str] = set()
    for record in records:
        code5 = record["municipality_code"]
        key = (record["prefecture"], record["municipality"])
        if key in canonical_by_name:
            fail(f"duplicate canonical municipality name for ordinary-grant join: {key}")
        canonical_by_name[key] = code5
        (ward_codes if is_special_ward(code5) else nonward_codes).add(code5)
    if len(ward_codes) != expected_wards:
        fail(f"Tokyo special ward count: expected {expected_wards}, got {len(ward_codes)}")
    if len(nonward_codes) != expected_count:
        fail(f"ordinary municipality count: expected {expected_count}, got {len(nonward_codes)}")

    parsed = parse_source(source, canonical_by_name, no_download=no_download)
    if len(parsed) != expected_count:
        missing = sorted(nonward_codes - set(parsed))
        extra = sorted(set(parsed) - nonward_codes)
        fail(
            f"ordinary-grant source municipality count: expected {expected_count}, got {len(parsed)}; "
            f"missing={missing[:10]} extra={extra[:10]}"
        )
    if set(parsed) != nonward_codes:
        missing = sorted(nonward_codes - set(parsed))
        extra = sorted(set(parsed) - nonward_codes)
        fail(f"ordinary-grant municipality join mismatch: missing={missing[:10]} extra={extra[:10]}")

    recipient_count = 0
    nonrecipient_count = 0
    capped_count = 0
    for record in records:
        code5 = record["municipality_code"]
        record["ordinary_grant_fiscal_year"] = int(source["fiscal_year"])
        record["ordinary_grant_source_stage"] = source["stage"]
        record["ordinary_grant_source_stage_label"] = source["stage_label"]
        if is_special_ward(code5):
            record["ordinary_grant_status"] = "special_ward_na"
            record["ordinary_grant_amount"] = None
            record["ordinary_grant_estimate"] = None
            record["balance_with_ordinary_grant_estimate"] = None
            record["ordinary_grant_source_row"] = None
            continue
        row = parsed[code5]
        amount = row["amount"]
        statutory_reference = float(record["tax_deduction_75pct_reference"])
        estimate = min(statutory_reference, amount)
        status = "recipient" if amount > 0 else "non_recipient"
        if status == "recipient":
            recipient_count += 1
        else:
            nonrecipient_count += 1
        if estimate < statutory_reference:
            capped_count += 1
        record["ordinary_grant_status"] = status
        record["ordinary_grant_amount"] = amount
        record["ordinary_grant_estimate"] = estimate
        record["balance_with_ordinary_grant_estimate"] = float(record["balance_before_tax_adjustment"]) + estimate
        record["ordinary_grant_source_row"] = row["source_row"]

    return {
        "ordinary_grant_source_row_count": len(parsed),
        "ordinary_grant_recipient_count": recipient_count,
        "ordinary_grant_nonrecipient_count": nonrecipient_count,
        "ordinary_grant_special_ward_count": len(ward_codes),
        "ordinary_grant_estimate_capped_count": capped_count,
    }


def validate_enriched_record(record: dict, *, assessment_year: int) -> None:
    if record["ordinary_grant_fiscal_year"] != assessment_year:
        fail(
            f"{record['municipality_code']} ordinary-grant fiscal year "
            f"{record['ordinary_grant_fiscal_year']} != tax assessment year {assessment_year}"
        )
    reference = float(record["tax_deduction_75pct_reference"])
    if is_special_ward(record["municipality_code"]):
        for field in ("ordinary_grant_amount", "ordinary_grant_estimate", "balance_with_ordinary_grant_estimate", "ordinary_grant_source_row"):
            if record[field] is not None:
                fail(f"{record['municipality_code']} special ward must have null {field}")
        if record["ordinary_grant_status"] != "special_ward_na":
            fail(f"{record['municipality_code']} special ward status mismatch")
        return
    amount = record["ordinary_grant_amount"]
    estimate = record["ordinary_grant_estimate"]
    if not isinstance(amount, int) or amount < 0:
        fail(f"{record['municipality_code']} invalid ordinary-grant amount {amount!r}")
    if not isinstance(estimate, (int, float)) or isinstance(estimate, bool) or not math.isfinite(float(estimate)):
        fail(f"{record['municipality_code']} invalid ordinary-grant estimate {estimate!r}")
    if estimate < 0 or estimate > reference + 1e-5 or estimate > amount + 1e-5:
        fail(f"{record['municipality_code']} ordinary-grant estimate outside bounds")
    expected = min(reference, amount)
    if not math.isclose(float(estimate), expected, abs_tol=1e-5):
        fail(f"{record['municipality_code']} ordinary-grant estimate formula mismatch")
    expected_balance = float(record["balance_before_tax_adjustment"]) + expected
    if not math.isclose(float(record["balance_with_ordinary_grant_estimate"]), expected_balance, abs_tol=1e-5):
        fail(f"{record['municipality_code']} adjusted fiscal-impact formula mismatch")
    expected_status = "recipient" if amount > 0 else "non_recipient"
    if record["ordinary_grant_status"] != expected_status:
        fail(f"{record['municipality_code']} ordinary-grant status mismatch")
    if not isinstance(record["ordinary_grant_source_row"], int) or record["ordinary_grant_source_row"] <= 0:
        fail(f"{record['municipality_code']} invalid ordinary-grant source row")
