#!/usr/bin/env python3
"""Build the offline dashboard data from official MIC workbooks.

The workbooks named in data/source_manifest.json are the only financial source
of truth.  This script verifies their SHA-256, parses the fixed official
columns, joins receipt and tax rows by municipality code (falling back to an
audited name join when an official code differs), validates every row, writes
an explicit normalized JSON file, and embeds a compact JS representation in
index.html.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path
from urllib.request import Request, urlopen

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
MANIFEST_PATH = ROOT / "data" / "source_manifest.json"
PROCESSED_PATH = ROOT / "data" / "processed" / "furusato_data.json"
RAW_DIR = ROOT / "data" / "raw"
INDEX_PATH = ROOT / "index.html"
MAX_AMOUNT = 10**14


def fail(message: str) -> "NoReturn":
    raise RuntimeError(message)


def atomic_write_text(path: Path, text: str) -> None:
    temporary = path.with_name(path.name + ".tmp")
    temporary.write_text(text, encoding="utf-8")
    temporary.replace(path)


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def normalize_text(value) -> str | None:
    if value is None:
        return None
    text = unicodedata.normalize("NFKC", str(value)).replace("\u3000", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text or None


def normalize_code(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if not math.isfinite(float(value)):
            return None
        if float(value).is_integer():
            digits = str(int(value))
        else:
            digits = str(value).split(".", 1)[0]
    else:
        digits = re.sub(r"\D", "", str(value))
    if not digits:
        return None
    digits = digits.zfill(6)
    return digits if len(digits) == 6 else None


def amount(value, *, field: str, row: int, allow_blank: bool = False) -> int | float:
    if value is None or (isinstance(value, str) and not value.strip()):
        if allow_blank:
            return 0
        fail(f"{field}: blank amount at source row {row}")
    if isinstance(value, bool):
        fail(f"{field}: boolean amount at source row {row}")
    try:
        dec = Decimal(str(value).replace(",", "").replace("円", "").strip())
    except (InvalidOperation, ValueError) as exc:
        fail(f"{field}: non-numeric amount {value!r} at source row {row}: {exc}")
    if not dec.is_finite():
        fail(f"{field}: non-finite amount at source row {row}")
    if dec < 0 or dec > MAX_AMOUNT:
        fail(f"{field}: out-of-range amount at source row {row}: {dec}")
    if dec == dec.to_integral_value():
        return int(dec)
    return float(dec)


def excel_col_number(value: str) -> int:
    result = 0
    for char in value.upper():
        if not "A" <= char <= "Z":
            fail(f"invalid Excel column: {value}")
        result = result * 26 + ord(char) - ord("A") + 1
    return result


def download_or_read(source: dict, key: str, *, no_download: bool) -> bytes:
    url_key = f"{key}_url"
    file_key = f"{key}_file"
    hash_key = f"{key}_sha256"
    path = RAW_DIR / source[file_key]
    data = path.read_bytes() if path.exists() else None
    if data is None and no_download:
        fail(f"missing cached source: {path}; run without --no-download")
    if data is None:
        path.parent.mkdir(parents=True, exist_ok=True)
        request = Request(
            source[url_key],
            headers={"User-Agent": "furusato-data-integrity/2.0 (+https://github.com/ryotamatsuki/furusatonozei)"},
        )
        with urlopen(request, timeout=120) as response:
            data = response.read()
        path.write_bytes(data)
    actual = sha256_bytes(data)
    expected = source[hash_key].lower()
    if actual != expected:
        fail(f"{key} SHA-256 mismatch for {source[url_key]}: expected {expected}, got {actual}")
    return data


def load_sheet(data: bytes, sheet_name: str):
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        fail(f"sheet {sheet_name!r} not found; available={workbook.sheetnames}")
    return workbook[sheet_name]


def parse_source_rows(source: dict, *, kind: str, data: bytes) -> dict[str, dict]:
    manifest_kind = "receipts" if kind == "receipt" else "tax"
    sheet_key = f"{manifest_kind}_sheet"
    start = int(source[f"{manifest_kind}_row_start"])
    columns = source[f"{manifest_kind}_columns"]
    sheet = load_sheet(data, source[sheet_key])
    code_index = None if columns.get("municipality_code") is None else excel_col_number(columns["municipality_code"]) - 1
    pref_index = excel_col_number(columns["prefecture"]) - 1
    name_index = excel_col_number(columns["municipality"]) - 1
    amount_fields = {
        "receipt": {"received": "received", "proxy": "proxy", "expense": "expense"},
        "tax": {
            "municipal_tax_deduction": "municipal_tax_deduction",
            "prefectural_tax_deduction": "prefectural_tax_deduction",
        },
    }
    indexes = {
        field: excel_col_number(columns[column]) - 1
        for field, column in amount_fields[kind].items()
    }
    rows: dict[str, dict] = {}
    duplicate_codes: list[tuple[str, int, int]] = []
    code_master: dict[tuple[str | None, str], str] = {}
    code_master_seen: dict[tuple[str | None, str], list[tuple[str, int]]] = defaultdict(list)
    code_sheet_key = f"{manifest_kind}_code_sheet"
    code_columns_key = f"{manifest_kind}_code_columns"
    if source.get(code_sheet_key):
        master_sheet = load_sheet(data, source[code_sheet_key])
        master_columns = source[code_columns_key]
        master_code_index = excel_col_number(master_columns["municipality_code"]) - 1
        master_pref_index = excel_col_number(master_columns["prefecture"]) - 1
        master_name_index = excel_col_number(master_columns["municipality"]) - 1
        for master_source_row, master_row in enumerate(master_sheet.iter_rows(values_only=True), start=1):
            if len(master_row) <= max(master_code_index, master_pref_index, master_name_index):
                continue
            master_code = normalize_code(master_row[master_code_index])
            master_pref = normalize_text(master_row[master_pref_index])
            master_name = normalize_text(master_row[master_name_index])
            if master_code and master_name and master_name not in {"0", "-", "—", "―"}:
                key = (master_pref, master_name)
                code_master_seen[key].append((master_code, master_source_row))
                code_master[key] = master_code
    duplicate_master_keys = {
        key: values for key, values in code_master_seen.items() if len(values) > 1
    }
    if duplicate_master_keys:
        fail(f"{kind} code master duplicate prefecture/name keys: {list(duplicate_master_keys.items())[:5]}")

    header_checks = source.get("_header_validation", {}).get(manifest_kind, {})
    for field, tokens in header_checks.items():
        column = columns.get(field)
        if not column:
            continue
        column_index = excel_col_number(column) - 1
        header_values = []
        for header_row in range(max(1, start - 5), start):
            value = sheet.cell(header_row, column_index + 1).value
            if value is not None:
                header_values.append(normalize_text(value) or "")
        header_text = " ".join(header_values)
        if not any(token in header_text for token in tokens):
            fail(
                f"{kind} header mismatch for {field} ({column}) in {source[sheet_key]}: "
                f"expected one of {tokens!r}, got {header_values!r}"
            )
    for source_row, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if source_row < start:
            continue
        required_indexes = [pref_index, name_index, *indexes.values()]
        if code_index is not None:
            required_indexes.append(code_index)
        if len(row) <= max(required_indexes):
            visible = [value for value in row[: max(pref_index, name_index) + 1] if value not in (None, "")]
            if visible:
                fail(f"{kind} short municipality row at source row {source_row}: {visible!r}")
            continue
        name = normalize_text(row[name_index])
        pref = normalize_text(row[pref_index])
        # Prefecture/aggregate rows in the MIC files have 0 or a blank in the
        # municipality-name column.  Only named six-digit municipality rows
        # are part of this dashboard.
        if not name or name in {"0", "-", "—", "―"}:
            continue
        if not pref:
            fail(f"{kind} municipality row has blank prefecture at source row {source_row}")
        raw_code6 = normalize_code(row[code_index]) if code_index is not None else None
        code6 = code_master.get((pref, name)) or raw_code6
        if kind == "receipt" and not code6:
            fail(f"receipt missing municipality code at source row {source_row}: {pref}{name}")
        code5 = code6[:5] if code6 else None
        row_key = code5 or f"{pref}|{name}"
        if row_key in rows:
            duplicate_codes.append((row_key, rows[row_key]["source_row"], source_row))
            continue
        parsed = {
            "code5": code5,
            "code6": code6,
            "raw_code6": raw_code6,
            "pref": pref,
            "name": name,
            "source_row": source_row,
        }
        if kind == "receipt":
            parsed["received"] = amount(row[indexes["received"]], field="received", row=source_row)
            parsed["expense"] = amount(row[indexes["expense"]], field="expense", row=source_row)
            parsed["proxy"] = amount(row[indexes["proxy"]], field="proxy", row=source_row, allow_blank=True)
        else:
            parsed["municipal_tax_deduction"] = amount(
                row[indexes["municipal_tax_deduction"]],
                field="municipal_tax_deduction",
                row=source_row,
            )
            parsed["prefectural_tax_deduction"] = amount(
                row[indexes["prefectural_tax_deduction"]],
                field="prefectural_tax_deduction",
                row=source_row,
            )
        rows[row_key] = parsed
    if duplicate_codes:
        fail(f"{kind} duplicate municipality codes: {duplicate_codes[:5]}")
    return rows


def join_year(year: int, source: dict, *, no_download: bool) -> tuple[list[dict], dict]:
    receipt_bytes = download_or_read(source, "receipts", no_download=no_download)
    tax_bytes = download_or_read(source, "tax", no_download=no_download)
    receipt_rows = parse_source_rows(source, kind="receipt", data=receipt_bytes)
    tax_rows = parse_source_rows(source, kind="tax", data=tax_bytes)
    expected_count = int(source["municipality_count"])
    if len(receipt_rows) != expected_count:
        fail(f"{year} receipt municipality count: expected {expected_count}, got {len(receipt_rows)}")
    if len(tax_rows) != expected_count:
        fail(f"{year} tax municipality count: expected {expected_count}, got {len(tax_rows)}")

    tax_by_code = {row["code5"]: row for row in tax_rows.values() if row["code5"]}
    tax_by_name: dict[tuple[str | None, str], list[dict]] = defaultdict(list)
    for row in tax_rows.values():
        tax_by_name[(row["pref"], row["name"])].append(row)

    def short_name(name: str) -> str:
        return re.sub(r"^.+?郡", "", name)

    tax_by_short_name: dict[tuple[str | None, str], list[dict]] = defaultdict(list)
    for row in tax_rows.values():
        tax_by_short_name[(row["pref"], short_name(row["name"]))].append(row)

    records: list[dict] = []
    corrections: list[dict] = []
    code_normalization_corrections: list[dict] = []
    for kind, source_rows in (("receipt", receipt_rows), ("tax", tax_rows)):
        for row in source_rows.values():
            if row.get("raw_code6") and row.get("code6") and row["raw_code6"] != row["code6"]:
                code_normalization_corrections.append(
                    {
                        "kind": kind,
                        "source_row": row["source_row"],
                        "prefecture": row["pref"],
                        "municipality": row["name"],
                        "raw_code6": row["raw_code6"],
                        "canonical_code6": row["code6"],
                    }
                )
    missing: list[str] = []
    allowed_name_mismatches = {
        (item["municipality_code"], item["receipt_name"], item["tax_name"]): item
        for item in source.get("allowed_code_name_mismatches", [])
    }
    for code5, receipt in receipt_rows.items():
        tax = tax_by_code.get(code5)
        if tax is not None and (tax["pref"], tax["name"]) != (receipt["pref"], receipt["name"]):
            allowed = allowed_name_mismatches.get((code5, receipt["name"], tax["name"]))
            if allowed is not None:
                corrections.append(
                    {
                        "receipt_code6": receipt["code6"],
                        "tax_code6": tax["code6"],
                        "receipt_raw_code6": receipt.get("raw_code6"),
                        "tax_raw_code6": tax.get("raw_code6"),
                        "receipt_name": receipt["name"],
                        "tax_name": tax["name"],
                        "reason": "manifest-allowed official code/name correction",
                    }
                )
            else:
                tax = None
        if tax is None:
            candidates = tax_by_name.get((receipt["pref"], receipt["name"]), [])
            if not candidates:
                candidates = tax_by_short_name.get((receipt["pref"], short_name(receipt["name"])), [])
            if len(candidates) != 1:
                fail(
                    f"{year} tax code/name join is not unique for "
                    f"{receipt['pref']}{receipt['name']} ({code5}): {len(candidates)} candidates"
                )
            tax = candidates[0]
            corrections.append(
                {
                    "receipt_code6": receipt["code6"],
                    "tax_code6": tax["code6"],
                    "receipt_raw_code6": receipt.get("raw_code6"),
                    "tax_raw_code6": tax.get("raw_code6"),
                    "receipt_name": receipt["name"],
                    "tax_name": tax["name"],
                    "reason": "official code/name join fallback",
                }
            )
        if tax is None:
            missing.append(f"{code5}:{receipt['pref']}{receipt['name']}")
            continue
        tax_value = tax["municipal_tax_deduction"]
        prefectural_tax_value = tax["prefectural_tax_deduction"]
        before = receipt["received"] - receipt["expense"] - receipt["proxy"] - tax_value
        grant75 = tax_value * 0.75
        reference = before + grant75
        records.append(
            {
                "receipt_fiscal_year": year,
                "receipt_fiscal_year_label": source["receipt_fiscal_year_label"],
                "tax_donation_calendar_year": int(source["tax_donation_calendar_year"]),
                "tax_assessment_fiscal_year": int(source["tax_assessment_fiscal_year"]),
                "tax_assessment_fiscal_year_label": source["tax_assessment_fiscal_year_label"],
                "municipality_code": code5,
                "municipality_code6": receipt["code6"],
                "prefecture": receipt["pref"],
                "municipality": receipt["name"],
                "received": receipt["received"],
                "expense": receipt["expense"],
                "proxy": receipt["proxy"],
                "municipal_tax_deduction": tax_value,
                "prefectural_tax_deduction": prefectural_tax_value,
                "resident_tax_deduction_total": tax_value + prefectural_tax_value,
                "balance_before_tax_adjustment": before,
                "tax_deduction_75pct_reference": grant75,
                "balance_with_75pct_reference": reference,
                "receipt_source_row": receipt["source_row"],
                "tax_source_row": tax["source_row"],
                "tax_source_code6": tax["code6"],
                "receipt_raw_code6": receipt.get("raw_code6"),
                "tax_raw_code6": tax.get("raw_code6"),
                "tax_source_prefecture": tax["pref"],
                "tax_source_municipality": tax["name"],
            }
        )
    if missing:
        fail(f"{year} tax join missing {len(missing)} rows: {missing[:5]}")
    if len(records) != expected_count:
        fail(f"{year} joined record count: expected {expected_count}, got {len(records)}")
    return records, {
        "receipt_source_row_count": len(receipt_rows),
        "tax_source_row_count": len(tax_rows),
        "tax_code_join_corrections": corrections,
        "code_normalization_corrections": code_normalization_corrections,
    }


def validate_manifest_corrections(year: int, diagnostics: dict, manifest: dict) -> None:
    """Require every observed raw-code/name correction to be documented."""
    actual = set()
    for item in diagnostics.get("code_normalization_corrections", []):
        actual.add(("code", item["kind"], item["prefecture"], item["municipality"], item["raw_code6"], item["canonical_code6"]))
    for item in diagnostics.get("tax_code_join_corrections", []):
        actual.add(("join", item["receipt_code6"], item["receipt_name"], item["tax_name"]))
    # The name-master entry is represented by the code-based join diagnostic;
    # normalize its signature to the same code/name tuple before comparison.
    normalized_expected = set()
    for item in manifest.get("join_corrections", {}).get(str(year), []):
        if item.get("raw_code6") and item.get("canonical_code6"):
            normalized_expected.add(("code", "receipt" if item["kind"].startswith("receipt_") else "tax", item.get("prefecture"), item.get("municipality"), item["raw_code6"], item["canonical_code6"]))
        else:
            receipt_code = item.get("receipt_code6") or (item.get("municipality_code") + "1")
            receipt_name = item.get("receipt_name") or item.get("canonical_name")
            tax_name = item.get("tax_name") or item.get("raw_name")
            normalized_expected.add(("join", receipt_code, receipt_name, tax_name))
    if actual != normalized_expected:
        fail(f"{year} undocumented or stale join corrections: actual={sorted(actual)!r}, manifest={sorted(normalized_expected)!r}")


def compact_number(value):
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def build_ui_record(record: dict, *, rank: int, reference_rank: int) -> dict:
    received = float(record["received"])
    expense = float(record["expense"])
    expense_rate = expense / received if received > 0 else None
    reference = record["balance_with_75pct_reference"]
    return {
        "code6": record["municipality_code6"],
        "code5": record["municipality_code"],
        "pref": record["prefecture"],
        "name": record["municipality"],
        "rank": rank,
        "received": compact_number(record["received"]),
        "expense": compact_number(record["expense"]),
        "proxy": compact_number(record["proxy"]),
        "netDonation": compact_number(record["received"] - record["expense"] - record["proxy"]),
        "taxDeduction": compact_number(record["municipal_tax_deduction"]),
        "prefecturalTaxDeduction": compact_number(record["prefectural_tax_deduction"]),
        "residentTaxDeductionTotal": compact_number(record["resident_tax_deduction_total"]),
        "grant75": compact_number(record["tax_deduction_75pct_reference"]),
        "beforeGrant": compact_number(record["balance_before_tax_adjustment"]),
        "afterGrant": compact_number(reference),
        "afterGrantStatutory": compact_number(reference),
        "referenceBalance": compact_number(reference),
        "statutoryGrant75": compact_number(record["tax_deduction_75pct_reference"]),
        "expenseRate": expense_rate,
        "balanceType": "参考値プラス" if reference >= 0 else "参考値マイナス",
        "receiptSourceRow": record["receipt_source_row"],
        "taxSourceRow": record["tax_source_row"],
        "taxSourceCode6": record["tax_source_code6"],
        "receiptRawSourceCode6": record["receipt_raw_code6"],
        "taxRawSourceCode6": record["tax_raw_code6"],
        "receiptFiscalYear": record["receipt_fiscal_year"],
        "taxDonationCalendarYear": record["tax_donation_calendar_year"],
        "taxAssessmentFiscalYear": record["tax_assessment_fiscal_year"],
        "rankStatutory": reference_rank,
    }


def build_embedded(normalized: dict) -> tuple[list[dict], dict, dict, list[str]]:
    years = normalized["years"]
    numeric_years = sorted((int(year) for year in years), key=int)
    latest_year = numeric_years[-1]
    latest_records = years[str(latest_year)]["records"]
    by_code = {row["municipality_code"]: row for row in latest_records}
    received_order = sorted(latest_records, key=lambda r: (-float(r["received"]), r["municipality_code"]))
    reference_order = sorted(latest_records, key=lambda r: (-float(r["balance_with_75pct_reference"]), r["municipality_code"]))
    received_rank = {row["municipality_code"]: index for index, row in enumerate(received_order, 1)}
    reference_rank = {row["municipality_code"]: index for index, row in enumerate(reference_order, 1)}
    # Keep the embedded order reproducible from the normalized sources alone.
    # Never derive it from an existing index.html, which would make a rebuild
    # depend on an unrelated prior artifact.
    data_order = [row["municipality_code"] for row in received_order]
    data = [build_ui_record(by_code[code], rank=received_rank[code], reference_rank=reference_rank[code]) for code in data_order]

    meta = {}
    history = {}
    for year in numeric_years:
        bucket = years[str(year)]
        source = bucket["source"]
        meta[str(year)] = {
            "label": source["receipt_fiscal_year_label"],
            "short": f"R{year - 2018}",
            "receiptFiscalYear": year,
            "taxDonationCalendarYear": source["tax_donation_calendar_year"],
            "taxAssessmentFiscalYear": source["tax_assessment_fiscal_year"],
            "taxAssessmentFiscalYearLabel": source["tax_assessment_fiscal_year_label"],
            "receiptUrl": source["receipts_url"],
            "receiptFile": source["receipts_file"],
            "receiptSheet": source["receipts_sheet"],
            "receiptColumns": source["receipts_columns"],
            "taxUrl": source["tax_url"],
            "taxFile": source["tax_file"],
            "taxSheet": source["tax_sheet"],
            "taxMunicipalColumn": source["tax_columns"]["municipal_tax_deduction"],
            "taxPrefecturalColumn": source["tax_columns"]["prefectural_tax_deduction"],
            "taxColumns": source["tax_columns"],
            "receiptSha256": source["receipts_sha256"],
            "taxSha256": source["tax_sha256"],
            "publicationDate": source["publication_date"],
            "municipalityCount": source["municipality_count"],
        }
        record_by_code = {row["municipality_code"]: row for row in bucket["records"]}
        history[str(year)] = [
            [
                code,
                record_by_code[code]["received"],
                record_by_code[code]["expense"],
                record_by_code[code]["proxy"],
                record_by_code[code]["municipal_tax_deduction"],
                record_by_code[code]["prefectural_tax_deduction"],
                record_by_code[code]["resident_tax_deduction_total"],
                record_by_code[code]["balance_with_75pct_reference"],
                record_by_code[code]["receipt_source_row"],
                record_by_code[code]["tax_source_row"],
                record_by_code[code]["receipt_raw_code6"],
                record_by_code[code]["tax_raw_code6"],
            ]
            for code in data_order
        ]
    return data, meta, history, data_order


def replace_const(text: str, name: str, value) -> str:
    encoded = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    pattern = rf"const {re.escape(name)} = .*?;\n"
    replacement = f"const {name} = {encoded};\n"
    updated, count = re.subn(pattern, replacement, text, count=1, flags=re.S)
    if count != 1:
        fail(f"index.html constant not found: {name}")
    return updated


def validate_normalized(normalized: dict) -> None:
    years = normalized.get("years", {})
    expected = list(map(str, range(normalized["period"]["start"], normalized["period"]["end"] + 1)))
    if list(years) != expected:
        fail(f"normalized years mismatch: expected {expected}, got {list(years)}")
    expected_count = normalized["period"]["municipality_count"]
    code_sets = []
    canonical_names: dict[str, tuple[str | None, str | None]] = {}
    for year in expected:
        records = years[year]["records"]
        if len(records) != expected_count:
            fail(f"{year} normalized count {len(records)} != {expected_count}")
        source = years[year]["source"]
        if source.get("retrieved_at") is None:
            fail(f"{year} source retrieved_at is missing")
        codes = [r["municipality_code"] for r in records]
        if len(set(codes)) != len(codes):
            fail(f"{year} normalized duplicate municipality code")
        code_sets.append(set(codes))
        receipt_rows = set()
        tax_rows = set()
        for r in records:
            if not re.fullmatch(r"\d{5}", r["municipality_code"]):
                fail(f"{year} invalid municipality code {r['municipality_code']}")
            for field in ("received", "expense", "proxy", "municipal_tax_deduction", "prefectural_tax_deduction", "resident_tax_deduction_total", "balance_before_tax_adjustment", "tax_deduction_75pct_reference", "balance_with_75pct_reference"):
                value = r[field]
                if not isinstance(value, (int, float)) or isinstance(value, bool) or not math.isfinite(float(value)):
                    fail(f"{year} {r['municipality_code']} invalid {field}: {value!r}")
                if abs(float(value)) > MAX_AMOUNT:
                    fail(f"{year} {r['municipality_code']} implausibly large {field}: {value!r}")
            if r["received"] < 0 or r["expense"] < 0 or r["proxy"] < 0 or r["municipal_tax_deduction"] < 0 or r["prefectural_tax_deduction"] < 0:
                fail(f"{year} {r['municipality_code']} negative amount")
            if r["receipt_fiscal_year"] != int(year) or r["receipt_fiscal_year_label"] != source["receipt_fiscal_year_label"]:
                fail(f"{year} {r['municipality_code']} receipt fiscal-year metadata")
            if r["tax_donation_calendar_year"] != int(source["tax_donation_calendar_year"]) or r["tax_assessment_fiscal_year"] != int(source["tax_assessment_fiscal_year"]):
                fail(f"{year} {r['municipality_code']} tax-period metadata")
            if r["receipt_source_row"] in receipt_rows or r["tax_source_row"] in tax_rows:
                fail(f"{year} {r['municipality_code']} source row duplicate")
            receipt_rows.add(r["receipt_source_row"])
            tax_rows.add(r["tax_source_row"])
            canonical = (r["prefecture"], r["municipality"])
            if r["municipality_code"] in canonical_names and canonical_names[r["municipality_code"]] != canonical:
                fail(f"{year} {r['municipality_code']} municipality name/prefecture differs between years")
            canonical_names[r["municipality_code"]] = canonical
            if not math.isclose(r["resident_tax_deduction_total"], r["municipal_tax_deduction"] + r["prefectural_tax_deduction"], abs_tol=1e-5):
                fail(f"{year} {r['municipality_code']} resident tax deduction total formula")
            if r["receipt_source_row"] <= 0 or r["tax_source_row"] <= 0:
                fail(f"{year} {r['municipality_code']} invalid source row")
            if not math.isclose(r["balance_before_tax_adjustment"], r["received"] - r["expense"] - r["proxy"] - r["municipal_tax_deduction"], abs_tol=1e-5):
                fail(f"{year} {r['municipality_code']} before balance formula")
            if not math.isclose(r["tax_deduction_75pct_reference"], r["municipal_tax_deduction"] * 0.75, abs_tol=1e-5):
                fail(f"{year} {r['municipality_code']} 75% formula")
            if not math.isclose(r["balance_with_75pct_reference"], r["balance_before_tax_adjustment"] + r["tax_deduction_75pct_reference"], abs_tol=1e-5):
                fail(f"{year} {r['municipality_code']} reference balance formula")
    if any(codes != code_sets[0] for codes in code_sets[1:]):
        fail("municipality code set differs between years")


def build(args: argparse.Namespace) -> None:
    manifest = json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))
    years = {}
    for year in range(manifest["period"]["start"], manifest["period"]["end"] + 1):
        source = dict(manifest["sources"][str(year)])
        source["_header_validation"] = manifest.get("header_validation", {})
        records, diagnostics = join_year(year, source, no_download=args.no_download)
        validate_manifest_corrections(year, diagnostics, manifest)
        source.pop("_header_validation", None)
        years[str(year)] = {"source": source, "diagnostics": diagnostics, "records": records}
    normalized = {
        "schema_version": 2,
        "generated_at": f"{manifest['dashboard']['data_cutoff_date']}T00:00:00+00:00",
        "manifest_schema_version": manifest["schema_version"],
        "period": manifest["period"],
        "source_page_url": manifest["source_page_url"],
        "latest_official_page_url": manifest["latest_official_page_url"],
        "latest_official_summary_url": manifest["latest_official_summary_url"],
        "normalization_notes": manifest.get("normalization_notes", {}),
        "years": years,
    }
    validate_normalized(normalized)
    PROCESSED_PATH.parent.mkdir(parents=True, exist_ok=True)
    atomic_write_text(PROCESSED_PATH, json.dumps(normalized, ensure_ascii=False, indent=2) + "\n")
    if args.no_index:
        print(f"built {PROCESSED_PATH.relative_to(ROOT)}")
        return
    data, meta, history, data_order = build_embedded(normalized)
    latest_source = manifest["sources"][str(manifest["period"]["end"])]
    latest_bucket = years[str(manifest["period"]["end"])]
    base_source_meta = {
        "auditVersion": f"{manifest['dashboard']['data_cutoff_date']}.1",
        "generatedAt": manifest["dashboard"]["data_cutoff_date"],
        "municipalityCount": manifest["period"]["municipality_count"],
        "receiptExpenseProxy": {
            "label": "総務省『ふるさと納税に関する現況調査』自治体別受入額等",
            "period": latest_source["receipt_fiscal_year_label"],
            "url": latest_source["receipts_url"],
            "file": latest_source["receipts_file"],
            "sha256": latest_source["receipts_sha256"],
            "sheet": latest_source["receipts_sheet"],
            "columns": latest_source["receipts_columns"],
        },
        "taxDeduction": {
            "label": "総務省『ふるさと納税に関する現況調査』市町村民税控除額",
            "period": latest_source["tax_assessment_fiscal_year_label"],
            "url": latest_source["tax_url"],
            "file": latest_source["tax_file"],
            "sha256": latest_source["tax_sha256"],
            "sheet": latest_source["tax_sheet"],
            "columns": latest_source["tax_columns"],
        },
        "taxCodeCorrections": latest_bucket["diagnostics"]["tax_code_join_corrections"],
        "periodDefinition": manifest["indicator_definitions"]["period_note"],
    }
    index = INDEX_PATH.read_text(encoding="utf-8")
    index = replace_const(index, "BASE_AMOUNT_SOURCE_META", base_source_meta)
    index = replace_const(index, "DATA", data)
    index = replace_const(index, "FIVE_YEAR_META", meta)
    index = replace_const(index, "FIVE_YEAR_HISTORY", history)
    index = replace_const(index, "FIVE_YEAR_FIELDS", ["received", "expense", "proxy", "taxDeduction", "prefecturalTaxDeduction", "residentTaxDeductionTotal", "referenceBalance", "receiptSourceRow", "taxSourceRow", "receiptRawSourceCode6", "taxRawSourceCode6"])
    atomic_write_text(INDEX_PATH, index)
    print(f"built {PROCESSED_PATH.relative_to(ROOT)}")
    print(f"embedded {len(data)} latest records and {len(history) * len(history[next(iter(history))])} historical records")
    print(f"latest fiscal year: {manifest['period']['end']}")
    correction_count = sum(len(bucket["diagnostics"]["tax_code_join_corrections"]) for bucket in years.values())
    print(f"tax code/name join fallbacks: {correction_count}")


def check_processed() -> None:
    if not PROCESSED_PATH.exists():
        fail(f"missing {PROCESSED_PATH}; run build_data.py first")
    normalized = json.loads(PROCESSED_PATH.read_text(encoding="utf-8"))
    validate_normalized(normalized)
    print(f"PROCESSED OK: {sum(len(bucket['records']) for bucket in normalized['years'].values())} records")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--no-download", action="store_true", help="use data/raw only")
    parser.add_argument("--no-index", action="store_true", help="write processed JSON but do not rewrite index.html")
    parser.add_argument("--check", action="store_true", help="validate processed JSON only")
    args = parser.parse_args()
    try:
        if args.check:
            check_processed()
        else:
            build(args)
    except Exception as exc:
        print(f"BUILD FAILED: {exc}", file=sys.stderr)
        raise


if __name__ == "__main__":
    main()
