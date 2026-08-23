#!/usr/bin/env python3
"""Inspect pinned candidate MIC ordinary-grant workbooks for parser design.

Temporary implementation diagnostic. Emits hashes, sheet metadata, and a small
row sample only; production parsing is implemented separately after schemas are
pinned.
"""
from __future__ import annotations

import hashlib
import io
import json
from pathlib import Path
from urllib.request import Request, urlopen

from openpyxl import load_workbook

SOURCES = {
    2021: {"stage": "initial", "url": "https://www.soumu.go.jp/main_content/000761570.xlsx"},
    2022: {"stage": "recalculated", "url": "https://www.soumu.go.jp/main_content/000850258.xlsx"},
    2023: {"stage": "recalculated", "url": "https://www.soumu.go.jp/main_content/000916157.xlsx"},
    2024: {"stage": "recalculated", "url": "https://www.soumu.go.jp/main_content/000983380.xlsx"},
    2025: {"stage": "recalculated", "url": "https://www.soumu.go.jp/main_content/001047530.xlsx"},
    2026: {"stage": "initial", "url": "https://www.soumu.go.jp/main_content/001083404.xlsx"},
}
OUTPUT = Path("/tmp/grant-workbook-inspection.json")


def download(url: str) -> bytes:
    req = Request(url, headers={"User-Agent": "furusato-grant-source-inspection/1.0"})
    with urlopen(req, timeout=120) as response:
        return response.read()


def value(v):
    if v is None or isinstance(v, (str, int, float, bool)):
        return v
    return str(v)


def inspect(year: int, source: dict) -> dict:
    data = download(source["url"])
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        samples = []
        for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
            vals = [value(v) for v in row[:20]]
            if any(v not in (None, "") for v in vals):
                samples.append({"row": r, "values": vals})
            if len(samples) >= 35:
                break
        sheets.append({
            "title": ws.title,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "samples": samples,
        })
    return {
        "fiscal_year": year,
        "stage": source["stage"],
        "url": source["url"],
        "sha256": hashlib.sha256(data).hexdigest(),
        "bytes": len(data),
        "sheets": sheets,
    }


def main() -> None:
    result = [inspect(year, source) for year, source in SOURCES.items()]
    OUTPUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    for item in result:
        print(item["fiscal_year"], item["stage"], item["sha256"], item["bytes"], [s["title"] for s in item["sheets"]])
    print(f"wrote {OUTPUT}")


if __name__ == "__main__":
    main()
