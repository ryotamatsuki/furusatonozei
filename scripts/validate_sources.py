#!/usr/bin/env python3
"""Independently reconcile every furusato source row with processed data.

The builder is intentionally not imported. Official XLSX files are read again
with a separate parser, their SHA-256 hashes are verified, and 22 source/raw
fields plus 3 derived fields are compared for every municipality-year record.
Identity joins use the source names preserved in the processed record; source
row numbers, raw codes and all monetary values are still read independently
from the official workbook.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import re
import sys
import unicodedata
from decimal import Decimal, InvalidOperation
from pathlib import Path
from urllib.request import Request, urlopen

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
MANIFEST = ROOT / "data" / "source_manifest.json"
PROCESSED = ROOT / "data" / "processed" / "furusato_data.json"
RAW_DIR = ROOT / "data" / "raw"
TOLERANCE = 1e-5
MAX_AMOUNT = 10**14

SOURCE_FIELDS = (
    "receipt_fiscal_year",
    "receipt_fiscal_year_label",
    "tax_donation_calendar_year",
    "tax_assessment_fiscal_year",
    "tax_assessment_fiscal_year_label",
    "municipality_code",
    "municipality_code6",
    "prefecture",
    "municipality",
    "received",
    "expense",
    "proxy",
    "municipal_tax_deduction",
    "prefectural_tax_deduction",
    "resident_tax_deduction_total",
    "receipt_source_row",
    "tax_source_row",
    "tax_source_code6",
    "receipt_raw_code6",
    "tax_raw_code6",
    "tax_source_prefecture",
    "tax_source_municipality",
)
DERIVED_FIELDS = (
    "balance_before_tax_adjustment",
    "tax_deduction_75pct_reference",
    "balance_with_75pct_reference",
)


def fail(message: str) -> "NoReturn":
    raise RuntimeError(message)


def norm(value) -> str | None:
    if value is None:
        return None
    text = unicodedata.normalize("NFKC", str(value)).replace("\u3000", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text or None


def code6(value) -> str | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if not math.isfinite(float(value)):
            return None
        digits = str(int(value)) if float(value).is_integer() else ""
    else:
        digits = re.sub(r"\D", "", str(value))
    if not digits:
        return None
    digits = digits.zfill(6)
    return digits if len(digits) == 6 else None


def amount(value, *, row: int, field: str, allow_blank: bool = False) -> int | float:
    if value is None or (isinstance(value, str) and not value.strip()):
        if allow_blank:
            return 0
        fail(f"{field}: blank at source row {row}")
    if isinstance(value, bool):
        fail(f"{field}: boolean at source row {row}")
    try:
        dec = Decimal(str(value).replace(",", "").replace("円", "").strip())
    except (InvalidOperation, ValueError) as exc:
        fail(f"{field}: non-numeric {value!r} at source row {row}: {exc}")
    if not dec.is_finite() or dec < 0 or dec > MAX_AMOUNT:
        fail(f"{field}: invalid {value!r} at source row {row}")
    return int(dec) if dec == dec.to_integral_value() else float(dec)


def col(label: str) -> int:
    value = 0
    for char in label.upper():
        if not "A" <= char <= "Z":
            fail(f"invalid Excel column {label}")
        value = value * 26 + ord(char) - 64
    return value - 1


def source_bytes(source: dict, kind: str, *, no_download: bool) -> bytes:
    prefix = "receipts" if kind == "receipt" else "tax"
    path = RAW_DIR / source[f"{prefix}_file"]
    if path.exists():
        data = path.read_bytes()
    elif no_download:
        fail(f"missing cached source {path}")
    else:
        path.parent.mkdir(parents=True, exist_ok=True)
        req = Request(source[f"{prefix}_url"], headers={"User-Agent": "furusato-independent-reconciliation/3.0"})
        with urlopen(req, timeout=120) as response:
            data = response.read()
        path.write_bytes(data)
    actual = hashlib.sha256(data).hexdigest()
    expected = source[f"{prefix}_sha256"].lower()
    if actual != expected:
        fail(f"{prefix} SHA-256 mismatch: expected {expected}, got {actual}")
    return data


def parse_sheet(source: dict, kind: str, expected: dict[tuple[str, str], dict], *, no_download: bool) -> dict[str, dict]:
    prefix = "receipts" if kind == "receipt" else "tax"
    data = source_bytes(source, kind, no_download=no_download)
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet_name = source[f"{prefix}_sheet"]
    if sheet_name not in wb.sheetnames:
        fail(f"{prefix} sheet {sheet_name!r} missing; available={wb.sheetnames}")
    ws = wb[sheet_name]
    columns = source[f"{prefix}_columns"]
    pref_i, name_i = col(columns["prefecture"]), col(columns["municipality"])
    raw_code_i = col(columns["municipality_code"]) if columns.get("municipality_code") else None
    if kind == "receipt":
        value_columns = {"received": col(columns["received"]), "proxy": col(columns["proxy"]), "expense": col(columns["expense"])}
    else:
        value_columns = {"municipal_tax_deduction": col(columns["municipal_tax_deduction"]), "prefectural_tax_deduction": col(columns["prefectural_tax_deduction"])}
    required = max(pref_i, name_i, *value_columns.values(), *([] if raw_code_i is None else [raw_code_i]))
    rows: dict[str, dict] = {}
    matched_keys: set[tuple[str, str]] = set()
    start = int(source[f"{prefix}_row_start"])
    for row_no, values in enumerate(ws.iter_rows(values_only=True), start=1):
        if row_no < start or len(values) <= required:
            continue
        pref, name = norm(values[pref_i]), norm(values[name_i])
        if not pref or not name:
            continue
        key = (pref, name)
        record = expected.get(key)
        if record is None:
            continue
        if key in matched_keys:
            fail(f"{kind}: duplicate source identity {pref}{name}")
        matched_keys.add(key)
        parsed = {
            "source_row": row_no,
            "prefecture": pref,
            "municipality": name,
            "raw_code6": code6(values[raw_code_i]) if raw_code_i is not None else None,
        }
        for field, index in value_columns.items():
            parsed[field] = amount(values[index], row=row_no, field=field, allow_blank=(field == "proxy"))
        rows[record["municipality_code"]] = parsed
    missing = sorted(set(r["municipality_code"] for r in expected.values()) - set(rows))
    if missing or len(rows) != len(expected):
        fail(f"{kind}: expected {len(expected)} municipality rows, got {len(rows)}; missing={missing[:10]}")
    return rows


def equal(a, b) -> bool:
    if isinstance(a, (int, float)) and not isinstance(a, bool) and isinstance(b, (int, float)) and not isinstance(b, bool):
        return math.isclose(float(a), float(b), rel_tol=0, abs_tol=TOLERANCE)
    return a == b


def audit(*, no_download: bool) -> tuple[list[dict], dict]:
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    processed = json.loads(PROCESSED.read_text(encoding="utf-8"))
    errors: list[dict] = []
    yearly: dict[str, dict] = {}

    for year in range(int(manifest["period"]["start"]), int(manifest["period"]["end"]) + 1):
        ys = str(year)
        source = manifest["sources"][ys]
        records = processed["years"][ys]["records"]
        if len(records) != int(source["municipality_count"]):
            fail(f"{year}: processed count {len(records)}")
        by_code = {r["municipality_code"]: r for r in records}
        if len(by_code) != len(records):
            fail(f"{year}: duplicate processed municipality code")
        receipt_expected = {(norm(r["prefecture"]), norm(r["municipality"])): r for r in records}
        tax_expected = {(norm(r["tax_source_prefecture"]), norm(r["tax_source_municipality"])): r for r in records}
        if len(receipt_expected) != len(records) or len(tax_expected) != len(records):
            fail(f"{year}: source identity is not unique")
        receipt_rows = parse_sheet(source, "receipt", receipt_expected, no_download=no_download)
        tax_rows = parse_sheet(source, "tax", tax_expected, no_download=no_download)
        tax_name_fallback_codes = {
            str(item["receipt_code6"])[:5]
            for item in manifest.get("join_corrections", {}).get(ys, [])
            if item.get("kind") == "tax_name_fallback"
        }

        year_errors_before = len(errors)
        for code, record in by_code.items():
            receipt, tax = receipt_rows[code], tax_rows[code]
            municipal = tax["municipal_tax_deduction"]
            prefectural = tax["prefectural_tax_deduction"]
            before = receipt["received"] - receipt["expense"] - receipt["proxy"] - municipal
            reference75 = municipal * 0.75
            official = {
                "receipt_fiscal_year": year,
                "receipt_fiscal_year_label": source["receipt_fiscal_year_label"],
                "tax_donation_calendar_year": int(source["tax_donation_calendar_year"]),
                "tax_assessment_fiscal_year": int(source["tax_assessment_fiscal_year"]),
                "tax_assessment_fiscal_year_label": source["tax_assessment_fiscal_year_label"],
                "municipality_code": code,
                "municipality_code6": record["municipality_code6"],
                "prefecture": receipt["prefecture"],
                "municipality": receipt["municipality"],
                "received": receipt["received"],
                "expense": receipt["expense"],
                "proxy": receipt["proxy"],
                "municipal_tax_deduction": municipal,
                "prefectural_tax_deduction": prefectural,
                "resident_tax_deduction_total": municipal + prefectural,
                "receipt_source_row": receipt["source_row"],
                "tax_source_row": tax["source_row"],
                "tax_source_code6": None if code in tax_name_fallback_codes else record["municipality_code6"],
                "receipt_raw_code6": receipt["raw_code6"],
                "tax_raw_code6": tax["raw_code6"],
                "tax_source_prefecture": tax["prefecture"],
                "tax_source_municipality": tax["municipality"],
                "balance_before_tax_adjustment": before,
                "tax_deduction_75pct_reference": reference75,
                "balance_with_75pct_reference": before + reference75,
            }
            for field in SOURCE_FIELDS + DERIVED_FIELDS:
                if not equal(official[field], record.get(field)):
                    errors.append({
                        "year": year,
                        "municipality_code": code,
                        "municipality": record["prefecture"] + record["municipality"],
                        "field": field,
                        "official_value": official[field],
                        "processed_value": record.get(field),
                    })
        yearly[ys] = {
            "records": len(records),
            "field_comparisons": len(records) * (len(SOURCE_FIELDS) + len(DERIVED_FIELDS)),
            "mismatches": len(errors) - year_errors_before,
        }

    requested_years = [str(y) for y in range(2020, 2025)]
    summary = {
        "yearly": yearly,
        "requested_2020_2024": {
            "records": sum(yearly[y]["records"] for y in requested_years),
            "field_comparisons": sum(yearly[y]["field_comparisons"] for y in requested_years),
            "mismatches": sum(yearly[y]["mismatches"] for y in requested_years),
        },
        "full_period": {
            "records": sum(v["records"] for v in yearly.values()),
            "field_comparisons": sum(v["field_comparisons"] for v in yearly.values()),
            "mismatches": len(errors),
        },
    }
    return errors, summary


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--no-download", action="store_true")
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    errors, summary = audit(no_download=args.no_download)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(errors, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    requested, full = summary["requested_2020_2024"], summary["full_period"]
    print("SOURCE RECONCILIATION " + ("OK" if not errors else "FAILED"))
    print(f"requested 2020-2024 records: {requested['records']}")
    print(f"requested field comparisons: {requested['field_comparisons']}")
    print(f"requested mismatches: {requested['mismatches']}")
    print(f"full records: {full['records']}")
    print(f"full field comparisons: {full['field_comparisons']}")
    print(f"full mismatches: {full['mismatches']}")
    if errors:
        print(json.dumps(errors[:100], ensure_ascii=False, indent=2), file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"SOURCE RECONCILIATION FAILED: {exc}", file=sys.stderr)
        raise
